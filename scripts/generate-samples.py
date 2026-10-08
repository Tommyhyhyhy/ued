from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
out = Path("public/samples")
out.mkdir(parents=True, exist_ok=True)
font = Path("C:/Windows/Fonts/arial.ttf")
if not font.exists():
    font = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
pdfmetrics.registerFont(TTFont("Vietnamese", str(font)))
c = canvas.Canvas(str(out / "sample.pdf"), pagesize=(595,842))
sections = [
 ("Học liệu minh họa UEDocs", ["Tài liệu được tạo riêng để trải nghiệm xem trước và tải xuống.", "Đây không phải giáo trình hoặc đề thi chính thức của UED.", "", "Học liệu được sẻ chia, tri thức được lan tỏa.", "", "01  Đọc có mục tiêu", "Trước khi bắt đầu, hãy viết ra điều bạn muốn hiểu.", "Tóm tắt mỗi phần bằng ngôn ngữ của chính mình."]),
 ("Từ kiến thức đến thực hành", ["02  Tự kiểm tra để ghi nhớ", "Đóng tài liệu và thử giải thích khái niệm vừa học.", "Dùng ví dụ gần gũi để liên hệ kiến thức với thực tế.", "", "03  Một lịch học vừa sức", "Chia nội dung thành những phần nhỏ, có thời gian nghỉ.", "Ôn lại sau một ngày, một tuần và một tháng."]),
 ("Sẻ chia có trách nhiệm", ["04  Ghi nguồn rõ ràng", "Chỉ đăng nội dung do bạn tạo hoặc có quyền chia sẻ.", "Không đưa dữ liệu cá nhân vào học liệu công khai.", "", "05  Cùng nhau học tốt hơn", "Để lại câu hỏi cụ thể và phản hồi với sự tôn trọng.", "", "Cảm ơn bạn đã góp phần xây dựng cộng đồng UEDocs."])
]
for i,(title,lines) in enumerate(sections):
    c.setFillColorRGB(.043,.122,.227)
    c.rect(0,735,595,107,fill=1,stroke=0)
    c.setFillColorRGB(1,1,1)
    c.setFont("Vietnamese",24)
    c.drawString(48,785,"UEDocs")
    c.setFont("Vietnamese",11)
    c.drawString(48,759,"HỌC LIỆU MINH HỌA • KHÔNG PHẢI TÀI LIỆU CHÍNH THỨC")
    c.setFillColorRGB(.09,.41,1)
    c.setFont("Vietnamese",21)
    c.drawString(48,686,title)
    c.setFillColorRGB(.15,.2,.3)
    c.setFont("Vietnamese",12)
    for j,line in enumerate(lines):
        c.drawString(48,633-j*29,line)
    c.setFillColorRGB(.5,.55,.65)
    c.setFont("Vietnamese",10)
    c.drawString(48,45,"UEDocs · Nội dung gốc do dự án tạo cho bản demo")
    c.drawRightString(547,45,f"{i+1} / 3")
    c.showPage()
c.save()
doc = Document()
doc.add_heading("Học liệu minh họa UEDocs",0)
doc.add_paragraph("Tài liệu mẫu do dự án tạo. Không phải tài liệu chính thức của trường.")
for title,lines in sections:
    doc.add_heading(title,1)
    for line in lines:
        if line: doc.add_paragraph(line)
doc.save(out/"sample.docx")
prs=Presentation()
for title,lines in sections:
    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text=title
    slide.placeholders[1].text="\n".join(line for line in lines if line)
prs.save(out/"sample.pptx")
wb=Workbook()
ws=wb.active
ws.title="Kế hoạch học tập"
for row in [["UEDocs - Dữ liệu minh họa"],["Ngày","Mục tiêu","Hoàn thành"],["Thứ hai","Đọc và tóm tắt khái niệm","Chưa"],["Thứ tư","Làm bài tập tự luyện","Chưa"],["Thứ sáu","Ôn tập và tự kiểm tra","Chưa"]]:
    ws.append(row)
ws.column_dimensions["A"].width=25
ws.column_dimensions["B"].width=45
ws.column_dimensions["C"].width=20
wb.save(out/"sample.xlsx")
with ZipFile(out/"sample.zip","w",ZIP_DEFLATED) as z:
    z.writestr("README.txt","UEDocs - Hoc lieu minh hoa, khong phai tai lieu chinh thuc.\n")
    z.write(out/"sample.pdf","hoc-lieu-minh-hoa.pdf")
print("Created PDF (3 pages), DOCX, PPTX, XLSX and ZIP demo fixtures.")
